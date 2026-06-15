# -*- coding: utf-8 -*-
"""Generate full Access Matrix Excel for Lake Cement MIS Dashboard."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Access Matrix"

# ── Colour palette ────────────────────────────────────────────────────────────
NAVY   = '173158'
ORANGE = 'F49545'
GREEN  = '239557'
WHITE  = 'FFFFFF'
L_GREY = 'F1F5F9'
M_GREY = 'E2E8F0'
D_GREY = '475569'
AMBER  = 'FEF3C7'
PURPLE = 'F3E8FF'
BLUE   = 'DBEAFE'
LGREEN = 'DCFCE7'
DGREY  = 'F8FAFC'

def fill(hex_):  return PatternFill('solid', fgColor=hex_)
def font(hex_, bold=False, sz=8): return Font(bold=bold, color=hex_, size=sz)
def center(): return Alignment(horizontal='center', vertical='center', wrap_text=True)
def left():   return Alignment(horizontal='left',   vertical='center', wrap_text=True)

thin  = Side(style='thin',   color=M_GREY)
thick = Side(style='medium', color=NAVY)
bdr   = Border(left=thin, right=thin, top=thin, bottom=thin)

# ── Access styling map ────────────────────────────────────────────────────────
STYLE = {
    'EDIT':   {'bg': LGREEN, 'fg': '166534', 'label': 'EDIT'},
    'VIEW':   {'bg': BLUE,   'fg': '1E40AF', 'label': 'VIEW'},
    'ENTRY':  {'bg': AMBER,  'fg': '92400E', 'label': 'ENTRY'},
    'OWN':    {'bg': 'FFF7ED','fg': 'B45309', 'label': 'OWN'},
    'PRINT':  {'bg': PURPLE, 'fg': '7C3AED', 'label': 'PRINT'},
    'ACCESS': {'bg': LGREEN, 'fg': '059669', 'label': 'ACCESS'},
    '-':      {'bg': DGREY,  'fg': 'CBD5E1', 'label': '—'},
}

def apply_cell(cell, code):
    s = STYLE.get(code, STYLE['-'])
    cell.value        = s['label']
    cell.fill         = fill(s['bg'])
    cell.font         = Font(bold=True, color=s['fg'], size=7.5)
    cell.alignment    = center()
    cell.border       = bdr

# ── Department / Level definitions (X-axis) ──────────────────────────────────
DEPTS = [
    ('CORPORATE / MD',   'CORP', ['MD / CEO',      'Sr. Management',   'Corp. Staff'  ]),
    ('HR & ADMIN',        'HR',   ['HR Manager',    'HR Supervisor',    'HR Officer'   ]),
    ('FINANCE',           'FIN',  ['Finance Manager','Finance Supvr.',  'Accounts Ofcr']),
    ('STORES',            'STR',  ['Stores Manager','Stores Supvr.',    'Stores Offcr' ]),
    ('POWER PLANT',       'CPP',  ['CPP HOD',       'CPP Supervisor',   'CPP Operator' ]),
    ('P & M',             'PM',   ['P&M HOD',       'P&M Supervisor',   'P&M Technician']),
    ('IT',                'IT',   ['Sys. Admin',     'IT Manager',       'IT Support'   ]),
    ('LOGISTICS',         'LOG',  ['Logistics HOD', 'Logistics Supvr.', 'Dispatch Ofcr']),
    ('MINING',            'MIN',  ['Mining Manager','Mining Supervisor','Mining Officer']),
    ('PRODUCTION & QC',   'PRD',  ['Production HOD','Prod. Supervisor', 'QC Operator'  ]),
]

# ── Screen / Function definitions (Y-axis) ───────────────────────────────────
# Each entry: ('__GROUP__', 'Group Name')  or  ('Screen Name', [30 values])
# 30 values = CORP×3, HR×3, FIN×3, STR×3, CPP×3, PM×3, IT×3, LOG×3, MIN×3, PRD×3

MATRIX = [
    ('__GROUP__', 'MAIN NAVIGATION'),
    ('Login / Authentication',
      ['ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS','ACCESS']),
    ('Monthly KPIs Dashboard  (index.html)',
      ['VIEW','VIEW','VIEW','VIEW','VIEW','-','VIEW','VIEW','VIEW','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-']),
    ('Daily Reports Hub  (daily-report.html)',
      ['VIEW','VIEW','VIEW','-','-','-','VIEW','VIEW','VIEW','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','VIEW']),
    ('Trends & Analytics  (analytics.html)',
      ['VIEW','VIEW','VIEW','VIEW','VIEW','-','VIEW','VIEW','VIEW','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-']),
    ('MIS Dashboard  (mis.html)',
      ['VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-']),
    ('Daily Entry Forms — Status Board',
      ['VIEW','VIEW','VIEW','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','VIEW']),

    ('__GROUP__', 'MONTHLY KPIs — DEPARTMENT VIEWS'),
    ('Executive Summary (All Depts)',
      ['VIEW','VIEW','VIEW','VIEW','VIEW','-','VIEW','VIEW','VIEW','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-']),
    ('E&I Department KPIs',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('Production Department KPIs',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','-']),
    ('Mines Department KPIs',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','-','-','-','-']),
    ('Mechanical Department KPIs',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('Power Plant Department KPIs',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('Packing Plant KPIs',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','-','-','-','-','VIEW','VIEW','-']),
    ('Stores Department KPIs',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('Mines Workshop KPIs',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-']),
    ('Overtime KPIs',
      ['VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-']),

    ('__GROUP__', 'MONTHLY KPIs — ACTIONS'),
    ('Fill / Edit Monthly KPI Data',
      ['ENTRY','-','-','ENTRY','ENTRY','-','-','-','-','ENTRY','ENTRY','-','ENTRY','ENTRY','-','ENTRY','ENTRY','-','-','-','-','ENTRY','ENTRY','-','ENTRY','ENTRY','-','ENTRY','ENTRY','-']),
    ('Export KPI Report  (Excel / PDF)',
      ['PRINT','PRINT','-','PRINT','PRINT','-','PRINT','PRINT','-','PRINT','VIEW','-','-','-','-','-','-','-','PRINT','-','-','-','-','-','-','-','-','-','-','-']),
    ('Print KPI Report',
      ['PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT']),

    ('__GROUP__', 'DAILY REPORTS — VIEW ACCESS'),
    ('Daily Production Report — View',
      ['VIEW','VIEW','VIEW','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','VIEW']),
    ('Power Plant Daily Report — View',
      ['VIEW','VIEW','VIEW','-','-','-','VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','VIEW']),
    ('Mines Daily Report — View',
      ['VIEW','VIEW','VIEW','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','VIEW']),

    ('__GROUP__', 'DAILY REPORTS — SUBMIT / EDIT'),
    ('Daily Production Report — Submit/Edit',
      ['EDIT','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY']),
    ('Power Plant Daily Report — Submit/Edit',
      ['EDIT','-','-','-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('Mines Daily Report — Submit/Edit',
      ['EDIT','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY','-','-','-']),

    ('__GROUP__', 'DAILY ENTRY FORMS — VIEW ACCESS'),
    ('Dispatch Form — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','-','-','-','-','VIEW','VIEW','VIEW']),
    ('Stoppage Form — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','VIEW']),
    ('Limestone Form — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','VIEW']),
    ('Quality Report Form — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','VIEW']),
    ('Pile Quality Form — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','VIEW']),
    ('Overtime Form — View',
      ['VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','VIEW']),
    ('Materials (Stores) Form — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','VIEW']),
    ('Power Consumption Form — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','VIEW']),
    ('Kiln Form — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','VIEW']),
    ('Cement Mill Form — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','VIEW']),
    ('CPP Form — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','VIEW']),

    ('__GROUP__', 'DAILY ENTRY FORMS — SUBMIT / EDIT'),
    ('Dispatch Form — Submit/Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY','-','-','-','-','-','-']),
    ('Stoppage Form — Submit/Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('Limestone Form — Submit/Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY','-','-','-']),
    ('Quality Report Form — Submit/Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY']),
    ('Pile Quality Form — Submit/Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY']),
    ('Overtime Form — Submit/Edit',
      ['-','-','-','EDIT','ENTRY','ENTRY','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('Materials (Stores) Form — Submit/Edit',
      ['-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('Power Consumption Form — Submit/Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('Kiln Form — Submit/Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY']),
    ('Cement Mill Form — Submit/Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY']),
    ('CPP Form — Submit/Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','EDIT','ENTRY','ENTRY','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),

    ('__GROUP__', 'MIS — HR TAB'),
    ('HR Tab — Access',
      ['VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','VIEW','-','-','-','-','-','-','-','-','-','-','-']),
    ('Workforce Strength — View',
      ['VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','VIEW','-','-','-','-','-','-','-','-','-','-','-']),
    ('Workforce Strength — Edit',
      ['EDIT','-','-','EDIT','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('Contractor Roster — View',
      ['VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','VIEW','-','-','-','-','-','-','-','-','-','-','-']),
    ('Contractor Roster — Edit',
      ['-','-','-','EDIT','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('Contractors — View',
      ['VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','VIEW','-','-','-','-','-','-','-','-','-','-','-']),
    ('Contractors — Edit',
      ['-','-','-','EDIT','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('CSR Activities — View',
      ['VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','VIEW','-','-','-','-','-','-','-','-','-','-','-']),
    ('CSR Activities — Edit / Add',
      ['-','-','-','EDIT','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),

    ('__GROUP__', 'MIS — ASSETS TAB'),
    ('Assets Tab — Access',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','-','-','-']),
    ('IT Assets — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','OWN','OWN','-','OWN','OWN','-','VIEW','VIEW','-','OWN','OWN','-','-','-','-','-','-','-']),
    ('IT Assets — Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','OWN','-','-','OWN','-','-','EDIT','EDIT','-','OWN','-','-','-','-','-','-','-','-']),
    ('Vehicles — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','OWN','OWN','-','OWN','OWN','-','VIEW','-','-','VIEW','VIEW','-','-','-','-','-','-','-']),
    ('Vehicles — Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','OWN','-','-','OWN','-','-','-','-','-','EDIT','OWN','-','-','-','-','-','-','-']),
    ('Spares (Assets) — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','OWN','OWN','-','OWN','OWN','-','VIEW','-','-','-','-','-','-','-','-','-','-','-']),
    ('Spares (Assets) — Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','OWN','-','-','OWN','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),
    ('CAPEX — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','OWN','OWN','-','OWN','OWN','-','OWN','-','-','-','-','-','-','-','-','-','-','-']),
    ('CAPEX — Edit / Approve',
      ['EDIT','-','-','-','-','-','EDIT','VIEW','-','-','-','-','OWN','-','-','OWN','-','-','OWN','-','-','-','-','-','-','-','-','-','-','-']),
    ('Land & Property — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','VIEW','-','-','-','-','-','-','-','-','-','-','-']),
    ('Land & Property — Edit',
      ['EDIT','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),

    ('__GROUP__', 'MIS — COMPLIANCE TAB'),
    ('Compliance Tab — Access',
      ['VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','-','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-']),
    ('Certificates — View',
      ['VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','OWN','OWN','-','OWN','OWN','-','VIEW','-','-','OWN','OWN','-','OWN','OWN','-','-','-','-']),
    ('Certificates — Edit / Upload',
      ['-','-','-','EDIT','VIEW','-','-','-','-','-','-','-','OWN','-','-','OWN','-','-','-','-','-','OWN','-','-','OWN','-','-','-','-','-']),
    ('Mining Permits & Licences — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','VIEW','-','-','-','-','-','VIEW','VIEW','-','-','-','-']),
    ('Mining Permits & Licences — Edit',
      ['-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','EDIT','VIEW','-','-','-','-']),
    ('Legal Cases — View',
      ['VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-','VIEW','-','-','OWN','OWN','-','OWN','OWN','-','-','-','-']),
    ('Legal Cases — Edit',
      ['-','-','-','EDIT','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','OWN','-','-','OWN','-','-','-','-','-']),
    ('Insurance — View',
      ['VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','OWN','OWN','-','OWN','OWN','-','VIEW','-','-','OWN','OWN','-','OWN','OWN','-','-','-','-']),
    ('Insurance — Edit / Add Claim',
      ['-','-','-','EDIT','VIEW','-','EDIT','VIEW','-','-','-','-','OWN','-','-','OWN','-','-','-','-','-','OWN','-','-','OWN','-','-','-','-','-']),

    ('__GROUP__', 'MIS — STORE TAB'),
    ('Store Tab — Access',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','VIEW','-','VIEW','-','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-']),
    ('Raw Materials — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-','-','-','-','-','-','-','VIEW','-','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-']),
    ('Raw Materials — Edit',
      ['-','-','-','-','-','-','-','-','-','EDIT','VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','EDIT','VIEW','-','EDIT','VIEW','-']),
    ('Spare Parts (Store) — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-','OWN','OWN','-','OWN','OWN','-','VIEW','-','-','-','-','-','OWN','OWN','-','OWN','OWN','-']),
    ('Spare Parts (Store) — Edit',
      ['-','-','-','-','-','-','-','-','-','EDIT','VIEW','-','OWN','-','-','OWN','-','-','-','-','-','-','-','-','OWN','-','-','OWN','-','-']),
    ('AMC Contracts — View',
      ['VIEW','VIEW','-','-','-','-','VIEW','VIEW','-','VIEW','VIEW','-','OWN','OWN','-','OWN','OWN','-','VIEW','-','-','-','-','-','-','-','-','-','-','-']),
    ('AMC Contracts — Edit',
      ['-','-','-','-','-','-','-','-','-','EDIT','VIEW','-','OWN','-','-','OWN','-','-','-','-','-','-','-','-','-','-','-','-','-','-']),

    ('__GROUP__', 'GLOBAL FUNCTIONS'),
    ('Export Reports  (Excel / PDF)',
      ['PRINT','PRINT','-','PRINT','PRINT','-','PRINT','PRINT','-','PRINT','VIEW','-','-','-','-','-','-','-','PRINT','-','-','-','-','-','-','-','-','-','-','-']),
    ('Manage Users  (User Administration)',
      ['-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','EDIT','-','-','-','-','-','-','-','-','-','-','-']),
    ('View Audit Log',
      ['VIEW','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','-','VIEW','VIEW','-','-','-','-','-','-','-','-','-','-']),
    ('Print Reports',
      ['PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT','PRINT']),
]

# ── Layout constants ──────────────────────────────────────────────────────────
ROW_LEGEND  = 1   # legend / title row
ROW_DEPT    = 2   # department group headers
ROW_LEVEL   = 3   # L1 / L2 / L3 labels
ROW_SUBLVL  = 4   # role name labels
DATA_START  = 5   # first data row

COL_SCREEN  = 1   # "Screen / Function" column
COL_DATA    = 2   # first data column (CORP-L1)

# ── Write column headers ──────────────────────────────────────────────────────
# Row 1: Title + legend
ws.merge_cells(start_row=ROW_LEGEND, start_column=COL_SCREEN, end_row=ROW_LEGEND, end_column=COL_SCREEN+1)
c = ws.cell(row=ROW_LEGEND, column=COL_SCREEN,
            value="Lake Cement Ltd — MIS Access Matrix  |  2026")
c.fill      = fill(NAVY)
c.font      = Font(bold=True, color=WHITE, size=11)
c.alignment = left()

# Legend chips in same row
legend_items = [
    ('ACCESS', LGREEN, '059669'), ('VIEW', BLUE, '1E40AF'),
    ('EDIT',  LGREEN,  '166534'), ('ENTRY', AMBER, '92400E'),
    ('OWN',   'FFF7ED','B45309'), ('PRINT', PURPLE,'7C3AED'),
    ('—  No Access', DGREY, 'CBD5E1'),
]
legend_col = COL_DATA + 2
for label, bg, fg in legend_items:
    lc = ws.cell(row=ROW_LEGEND, column=legend_col, value=label)
    lc.fill      = fill(bg)
    lc.font      = Font(bold=True, color=fg, size=7.5)
    lc.alignment = center()
    lc.border    = bdr
    legend_col  += 1

# Screen column header
for r in (ROW_DEPT, ROW_LEVEL, ROW_SUBLVL):
    c = ws.cell(row=r, column=COL_SCREEN)
    c.fill = fill(NAVY)
    c.font = Font(bold=True, color=WHITE, size=9)
    c.alignment = center()
ws.merge_cells(start_row=ROW_DEPT,  start_column=COL_SCREEN, end_row=ROW_SUBLVL, end_column=COL_SCREEN)
ws.cell(row=ROW_DEPT, column=COL_SCREEN).value = "SCREEN / FUNCTION"

# Department group headers (Row 2), Level headers (Row 3), Sub-level (Row 4)
col = COL_DATA
DEPT_COL_STARTS = {}
for dept_full, dept_code, levels in DEPTS:
    DEPT_COL_STARTS[dept_code] = col
    # Merge 3 columns for dept header
    ws.merge_cells(start_row=ROW_DEPT, start_column=col, end_row=ROW_DEPT, end_column=col+2)
    dc = ws.cell(row=ROW_DEPT, column=col, value=dept_full)
    dc.fill      = fill(NAVY)
    dc.font      = Font(bold=True, color=WHITE, size=9)
    dc.alignment = center()
    dc.border    = Border(left=thick, right=thick, top=thin, bottom=thin)

    for i, (lv, sublv) in enumerate(zip(['L1','L2','L3'], levels)):
        lc = ws.cell(row=ROW_LEVEL, column=col+i, value=lv)
        lc.fill      = fill('1E3A5F')
        lc.font      = Font(bold=True, color=WHITE, size=8)
        lc.alignment = center()
        lc.border    = bdr

        sc = ws.cell(row=ROW_SUBLVL, column=col+i, value=sublv)
        sc.fill      = fill('2D4E7A')
        sc.font      = Font(color=WHITE, size=7)
        sc.alignment = center()
        sc.border    = bdr
    col += 3

TOTAL_COLS = col - 1  # last data column index

# ── Write data rows ───────────────────────────────────────────────────────────
row = DATA_START
for entry in MATRIX:
    if entry[0] == '__GROUP__':
        # Group header row — span entire width
        group_name = entry[1]
        ws.merge_cells(start_row=row, start_column=COL_SCREEN,
                       end_row=row, end_column=TOTAL_COLS)
        gc = ws.cell(row=row, column=COL_SCREEN, value=group_name)
        gc.fill      = fill(ORANGE)
        gc.font      = Font(bold=True, color=WHITE, size=9)
        gc.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        gc.border    = Border(left=thick, right=thick, top=thick, bottom=thick)
        ws.row_dimensions[row].height = 16
        row += 1
    else:
        screen_name, values = entry
        # Screen label
        sc = ws.cell(row=row, column=COL_SCREEN, value=screen_name)
        sc.fill      = fill(L_GREY)
        sc.font      = Font(size=8.5, color='1E293B')
        sc.alignment = Alignment(horizontal='left', vertical='center',
                                 indent=1, wrap_text=False)
        sc.border    = Border(left=thick, right=thin, top=thin, bottom=thin)
        # Access cells
        for ci, code in enumerate(values):
            cell = ws.cell(row=row, column=COL_DATA + ci)
            apply_cell(cell, code)
        ws.row_dimensions[row].height = 14
        row += 1

# ── Column widths ─────────────────────────────────────────────────────────────
ws.column_dimensions[get_column_letter(COL_SCREEN)].width = 34
for ci in range(len(DEPTS) * 3):
    ws.column_dimensions[get_column_letter(COL_DATA + ci)].width = 8.5

# ── Freeze panes ──────────────────────────────────────────────────────────────
ws.freeze_panes = get_column_letter(COL_DATA) + str(DATA_START)

# ── Tab colour ────────────────────────────────────────────────────────────────
ws.sheet_properties.tabColor = ORANGE

# ── Save ──────────────────────────────────────────────────────────────────────
out_path = r'C:\Users\USER\LakeCementDashboard\LCL_Access_Matrix_2026.xlsx'
wb.save(out_path)
print("Saved:", out_path)
