#!/usr/bin/env python3
"""
extract-production.py — reads the plant's Master Production Report
workbooks (April/May/June 2026) and writes one JSON file per month of
{date: {fact: value}}, ready for import-production.js to POST.

Why Python/openpyxl and not the Node importer's own xlsx parsing:
ExcelJS was found to silently drop the cached formula RESULT for a
handful of cells in these workbooks (verified: 2 of 30 days on one
column alone) while still parsing the cell as a formula object — so a
Node-only extractor would produce quietly incomplete data on an
unpredictable subset of days. openpyxl's data_only=True read has been
cross-checked cell-by-cell against known values throughout this project
and has not shown the same gap, so extraction lives here; import-
production.js's job is reduced to reading this JSON and POSTing it.

Column positions are NOT stable across month files (April's Gypsum
block has one "Consumption" column where May/June split "Cons dry"/
"Cons wet", which shifts every later column by one) — so every column
is located by scanning row-1 group labels and row-2 sub-labels fresh
per workbook, not by hardcoded letters. See resolve_cols().

Usage: python scripts/extract-production.py
Writes: scripts/_production_<month>.json (one per configured month)
"""
import json
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))

MONTHS = [
    {"file": os.path.join(ROOT, "full Production Report april and may", "Master production report APRIL 26'.xlsx"),
     "year": 2026, "month": 4, "days": 30},
    {"file": os.path.join(ROOT, "full Production Report april and may", "Master production report MAY 26'.xlsx"),
     "year": 2026, "month": 5, "days": 31},
    {"file": os.path.join(ROOT, "Master production report JUNE 26'.xlsx"),
     "year": 2026, "month": 6, "days": 30},
]


def norm(v):
    return str(v).strip().upper() if v is not None else ""


def find_col(ws, row_num, label, from_col=1, to_col=250):
    target = norm(label)
    for c in range(from_col, to_col + 1):
        if norm(ws.cell(row=row_num, column=c).value) == target:
            return c
    return None


def find_group_col(ws, row1_label, row2_label, width=6):
    target1 = norm(row1_label)
    for c in range(1, 251):
        if norm(ws.cell(row=1, column=c).value) == target1:
            found = find_col(ws, 2, row2_label, c, c + width - 1)
            if found:
                return found
    return None


def find_group_col_any(ws, row1_label, row2_labels, width=6):
    for lbl in row2_labels:
        c = find_group_col(ws, row1_label, lbl, width)
        if c:
            return c
    return None


def resolve_cols(p2, cl, ce):
    cols = {}
    cols['kiln_run_hrs'] = find_group_col(p2, 'Kiln', 'Rng Hrs')
    cols['cm1_425r_mt'] = find_group_col(p2, 'Cement Mill (42.5)', 'Production')
    cols['cm1_325n_mt'] = find_group_col(p2, 'Cement Mill (32.5', 'Production')
    cols['cm1_425n_mt'] = find_group_col(p2, 'Cement Mill (42.5N)', 'Production')
    cols['cm1_opc_mt'] = find_group_col(p2, 'Cement Mill (OPC)', 'Production')
    cols['cm2_425r_mt'] = find_group_col(p2, 'Cement Mill- 2 (42.5)', 'Production')
    cols['cm2_325n_mt'] = find_group_col(p2, 'Cement Mill -2 (32.5', 'Production')
    cols['cm2_425n_mt'] = find_group_col(p2, 'Cement Mill -2(42.5N)', 'Production')
    cols['cm2_425n_bm_mt'] = find_group_col(p2, 'Cement Mill -2(42.5 N BM)', 'Production')

    cols['kiln_clinker_open'] = find_group_col(cl, 'Clinker', 'Opening Stock')
    cols['kiln_prod_mt'] = find_group_col(cl, 'Clinker', 'Production')
    cols['cm_clinker_own_cons'] = find_group_col(cl, 'Clinker', 'Consumption')
    cols['coal_kiln'] = find_group_col(cl, 'KILN COAL', 'Consumption')
    cols['coal_cpp'] = find_group_col(cl, 'cpp coal', 'Consumption')
    purch_recv = find_col(cl, 2, 'CLINKER PURCHASED')
    cols['cm_clinker_purch_open'] = (purch_recv - 1) if purch_recv else None
    cols['cm_clinker_purch_recv'] = purch_recv
    cols['cm_clinker_purch_cons'] = (purch_recv + 1) if purch_recv else None

    cols['cm_gypsum_open'] = find_group_col(ce, 'Gypsum', 'Opening Stock')
    cols['cm_gypsum_recv'] = find_group_col(ce, 'Gypsum', 'Production')
    cols['cm_gypsum_cons'] = find_group_col_any(ce, 'Gypsum', ['Cons wet', 'Consumption'])
    cols['starex_open'] = find_group_col(ce, 'Starex', 'Opening Stock')
    cols['starex_recv'] = find_group_col(ce, 'Starex', 'Production')
    cols['starex_cons'] = find_group_col(ce, 'Starex', 'Consumption')
    cols['mapei_open'] = find_group_col(ce, 'Mapei', 'Opening Stock')
    cols['mapei_recv'] = find_group_col(ce, 'Mapei', 'Production')
    cols['mapei_cons'] = find_group_col(ce, 'Mapei', 'Consumption')
    return cols


def num(ws, row, col):
    if not col:
        return None
    v = ws.cell(row=row, column=col).value
    return v if isinstance(v, (int, float)) else None


def r2(n):
    return round(n * 100) / 100


def main():
    for m in MONTHS:
        if not os.path.exists(m['file']):
            print(f"(skipping {m['file']} — not found)")
            continue
        wb = openpyxl.load_workbook(m['file'], data_only=True)
        p2, cl, ce = wb['Production2'], wb['Clinker'], wb['Cement']
        cols = resolve_cols(p2, cl, ce)
        missing = [k for k, v in cols.items() if not v]
        print(f"=== {os.path.basename(m['file'])} ({m['year']}-{m['month']:02d}, {m['days']} days) ===")
        if missing:
            print(f"  WARNING columns not found (omitted): {missing}")

        out = {}
        for day in range(1, m['days'] + 1):
            row = day + 2
            data = {}

            def setf(fact, col):
                n = num(p2 if fact.startswith('kiln_run') or fact.startswith('cm1') or fact.startswith('cm2') else
                        (cl if fact in ('kiln_clinker_open', 'kiln_prod_mt', 'cm_clinker_own_cons', 'coal_kiln',
                                        'coal_cpp', 'cm_clinker_purch_open', 'cm_clinker_purch_recv',
                                        'cm_clinker_purch_cons') else ce),
                        row, col)
                if n is not None:
                    data[fact] = r2(n)

            for fact in ['kiln_run_hrs', 'cm1_425r_mt', 'cm1_325n_mt', 'cm1_425n_mt', 'cm1_opc_mt',
                         'cm2_425r_mt', 'cm2_325n_mt', 'cm2_425n_mt', 'cm2_425n_bm_mt',
                         'kiln_clinker_open', 'kiln_prod_mt', 'cm_clinker_own_cons', 'coal_kiln', 'coal_cpp',
                         'cm_clinker_purch_open', 'cm_clinker_purch_recv', 'cm_clinker_purch_cons',
                         'cm_gypsum_open', 'cm_gypsum_recv', 'cm_gypsum_cons']:
                setf(fact, cols.get(fact))

            starex_open = num(ce, row, cols['starex_open'])
            starex_recv = num(ce, row, cols['starex_recv'])
            starex_cons = num(ce, row, cols['starex_cons'])
            mapei_open = num(ce, row, cols['mapei_open'])
            mapei_recv = num(ce, row, cols['mapei_recv'])
            mapei_cons = num(ce, row, cols['mapei_cons'])
            if starex_open is not None or mapei_open is not None:
                data['cm_grind_aid_open'] = r2((starex_open or 0) + (mapei_open or 0))
            if starex_recv is not None or mapei_recv is not None:
                data['cm_grind_aid_recv'] = r2((starex_recv or 0) + (mapei_recv or 0))
            if starex_cons is not None or mapei_cons is not None:
                data['cm_grind_aid_cons'] = r2((starex_cons or 0) + (mapei_cons or 0))

            if any(abs(v) > 0.001 for v in data.values() if isinstance(v, (int, float))):
                date_str = f"{m['year']}-{m['month']:02d}-{day:02d}"
                out[date_str] = data

        out_path = os.path.join(HERE, f"_production_{m['year']}-{m['month']:02d}.json")
        with open(out_path, 'w', encoding='utf-8') as f:
            json.dump(out, f, indent=1)
        print(f"  wrote {len(out)} days -> {out_path}\n")


if __name__ == '__main__':
    main()
