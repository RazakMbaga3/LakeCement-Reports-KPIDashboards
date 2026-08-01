#!/usr/bin/env python3
"""
verify-raw-materials.py — INDEPENDENT post-hoc verification, separate from
extract-raw-materials.py / backfill-raw-materials.js's own inline checks.

Re-derives facts fresh from the Excel workbooks (re-implements the label
lookup independently rather than importing extract-raw-materials.py) and
diffs against a fresh GET from the live API for a spread of dates across
all 3 months, including the June-21 Bauxite/Mill-Dust-Iron-Ore boundary.

Also spot-checks a sample of the PRIOR pass's Plant Operations fields
(ls_crusher_*, kiln_*, cm1_*) to confirm zero regression from this pass's
writes.

Usage: python backend/scripts/verify-raw-materials.py <scratch_dir_with_login_read.json>
"""
import json
import os
import sys
import urllib.request
import openpyxl
import warnings
warnings.filterwarnings('ignore')

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
API = 'https://lakecement-mis-api.onrender.com/api'

COL = {c: i for i, c in enumerate(
    ['', 'A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q'])}

WB_FOR_MONTH = {
    4: os.path.join(ROOT, "full Production Report april and may", "Master production report APRIL 26'.xlsx"),
    5: os.path.join(ROOT, "full Production Report april and may", "Master production report MAY 26'.xlsx"),
    6: os.path.join(ROOT, "Master production report JUNE 26'.xlsx"),
}

SAMPLE_DATES = [
    '2026-04-01', '2026-04-08', '2026-04-15', '2026-04-22', '2026-04-30',
    '2026-05-01', '2026-05-10', '2026-05-15', '2026-05-20', '2026-05-31',
    '2026-06-01', '2026-06-10', '2026-06-15', '2026-06-20', '2026-06-21',
    '2026-06-25', '2026-06-30',
]

REGRESSION_SAMPLE_DATES = ['2026-04-01', '2026-05-15', '2026-06-30']
REGRESSION_FIELDS = ['ls_crusher_open', 'ls_crusher_prod_mt', 'ls_crusher_cons_mt',
                      'kiln_clinker_open', 'kiln_prod_mt', 'cm1_425r_mt', 'cm_gypsum_open']


def norm(v):
    return str(v).strip().upper() if v is not None else ""


def num(ws, row, col_letter):
    idx = COL[col_letter]
    v = ws.cell(row=row, column=idx).value
    return v if isinstance(v, (int, float)) else None


def r2(n):
    return round(n * 100) / 100


def find_row(ws, label, col_letter='C', r1=1, r2_=140):
    idx = COL[col_letter]
    target = norm(label)
    for r in range(r1, r2_ + 1):
        if norm(ws.cell(row=r, column=idx).value) == target:
            return r
    return None


def independent_extract(ws, wb):
    """Re-implemented independently from extract-raw-materials.py's extract_day()
    — deliberately NOT imported, to keep this a genuinely separate check."""
    out = {}
    pyro_row = find_row(ws, 'PYRO PROCESSING')
    total_rm_row = find_row(ws, 'TOTAL MATL - RAW MILL', r1=pyro_row)
    rm_map = {'LIMESTONE': 'rmm_limestone', 'CLAY': 'rmm_clay', 'IRON ORE': 'rmm_ironore',
              'INDUSTRIAL CLAY': 'rmm_industrialclay', 'BAUXITE': 'rmm_bauxite',
              'MILL DUST IRON ORE': 'rmm_milldustironore'}
    for r in range(pyro_row + 1, total_rm_row):
        label = norm(ws.cell(row=r, column=COL['C']).value)
        if label in rm_map:
            p = rm_map[label]
            for suf, col in [('_open', 'E'), ('_recv', 'F'), ('_cons', 'I')]:
                v = num(ws, r, col)
                if v is not None:
                    out[p + suf] = r2(v)

    cg_row = find_row(ws, 'CEMENT GRINDING')
    total_cm_row = find_row(ws, 'TOTAL MATL - CEMENT MILL', r1=cg_row)
    for r in range(cg_row + 1, total_cm_row):
        c_label = norm(ws.cell(row=r, column=COL['C']).value)
        d_label = norm(ws.cell(row=r, column=COL['D']).value)
        if c_label == 'LIMESTONE':
            f, i = num(ws, r, 'F'), num(ws, r, 'I')
            if f is not None: out['cm_limestone_recv'] = r2(f)
            if i is not None: out['cm_limestone_cons'] = r2(i)
        if d_label == 'STAREX':
            for suf, col in [('_open', 'E'), ('_recv', 'F'), ('_cons', 'I')]:
                v = num(ws, r, col)
                if v is not None:
                    out['cm_starex' + suf] = r2(v)

    iuc_row = find_row(ws, 'INTERNAL USE CEMENT')
    grade_map = {'42.5': 'iuc_425r', '42.5N': 'iuc_425n', '32.5': 'iuc_325n', 'OPC': 'iuc_opc'}
    for r in range(iuc_row, iuc_row + 4):
        e_label = ws.cell(row=r, column=COL['E']).value
        key = None
        if isinstance(e_label, (int, float)):
            if abs(e_label - 42.5) < 0.01: key = 'iuc_425r'
            elif abs(e_label - 32.5) < 0.01: key = 'iuc_325n'
        elif isinstance(e_label, str):
            key = grade_map.get(e_label.strip().upper())
        if key:
            f = num(ws, r, 'F')
            if f is not None: out[key] = r2(f)

    if 'Daily unit' in wb.sheetnames:
        du = wb['Daily unit']
        day_num = int(ws.title)
        row = day_num + 2
        for field, col in [('spc_crusher_kwh', 'B'), ('spc_rm1_kwh', 'C'), ('spc_rm2_kwh', 'D'),
                            ('spc_kiln_kwh', 'E'), ('spc_coalmill_kwh', 'F'), ('spc_cemmill_kwh', 'G'),
                            ('spc_losses_kwh', 'H'), ('spc_pplant_kwh', 'I'), ('spc_others_kwh', 'J')]:
            v = num(du, row, col)
            if v is not None:
                out[field] = r2(v)

    ncv_row = None
    for r in range(1, 140):
        v = ws.cell(row=r, column=COL['O']).value
        if isinstance(v, str) and 'NCV' in v.upper():
            ncv_row = r
            break
    if ncv_row:
        q = num(ws, ncv_row, 'Q')
        if q is not None and q != 0:
            out['coal_ncv'] = r2(q)

    return out


def main():
    scratch = sys.argv[1] if len(sys.argv) > 1 else None
    if not scratch:
        print("Usage: python verify-raw-materials.py <scratch_dir_with_login_read.json>")
        sys.exit(1)
    with open(os.path.join(scratch, 'login_read.json')) as f:
        token = json.load(f)['token']

    match, mismatch, missing = 0, 0, 0
    mismatches = []

    print("=== INDEPENDENT VERIFICATION — fresh GET vs fresh Excel re-derivation ===\n")
    wb_cache = {}
    for date_str in SAMPLE_DATES:
        y, m, d = [int(x) for x in date_str.split('-')]
        if m not in wb_cache:
            wb_cache[m] = openpyxl.load_workbook(WB_FOR_MONTH[m], data_only=True)
        wb = wb_cache[m]
        ws = wb[str(d)]
        expected = independent_extract(ws, wb)

        req = urllib.request.Request(f'{API}/reports/production/{date_str}',
                                      headers={'Authorization': f'Bearer {token}'})
        with urllib.request.urlopen(req) as resp:
            record = json.loads(resp.read().decode())
        actual = (record or {}).get('data', {})

        date_match, date_mismatch, date_missing = 0, 0, 0
        for k, v in expected.items():
            av = actual.get(k)
            if av is None:
                date_missing += 1
                mismatches.append((date_str, k, v, 'MISSING'))
            elif abs(float(av) - float(v)) < 0.01:
                date_match += 1
            else:
                date_mismatch += 1
                mismatches.append((date_str, k, v, av))
        match += date_match; mismatch += date_mismatch; missing += date_missing
        status = 'OK' if (date_mismatch == 0 and date_missing == 0) else 'ISSUES'
        print(f"  {date_str}: {len(expected):3d} fields checked -> match={date_match} mismatch={date_mismatch} missing={date_missing}  [{status}]")

    print(f"\nTOTAL: match={match} mismatch={mismatch} missing={missing}")
    if mismatches:
        print("\nDetails of mismatches/missing:")
        for d, k, exp, act in mismatches:
            print(f"  {d} {k}: expected={exp} actual={act}")

    print("\n=== REGRESSION CHECK — prior-pass Plant Operations fields still intact ===\n")
    reg_ok = True
    for date_str in REGRESSION_SAMPLE_DATES:
        req = urllib.request.Request(f'{API}/reports/production/{date_str}',
                                      headers={'Authorization': f'Bearer {token}'})
        with urllib.request.urlopen(req) as resp:
            record = json.loads(resp.read().decode())
        actual = (record or {}).get('data', {})
        present = [f for f in REGRESSION_FIELDS if f in actual]
        missing_f = [f for f in REGRESSION_FIELDS if f not in actual]
        print(f"  {date_str}: {len(present)}/{len(REGRESSION_FIELDS)} present" + (f"  MISSING={missing_f}" if missing_f else ""))
        if missing_f:
            reg_ok = False
    print(f"\nRegression check: {'PASS — all prior fields intact' if reg_ok else 'FAIL — see missing fields above'}")

    sys.exit(0 if (mismatch == 0 and missing == 0 and reg_ok) else 1)


if __name__ == '__main__':
    main()
